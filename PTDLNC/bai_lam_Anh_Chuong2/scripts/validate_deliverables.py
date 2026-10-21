"""Independent checks of delivered data, analysis, report evidence, and execution."""
from pathlib import Path
import hashlib
import json
import re
import numpy as np
import pandas as pd
import nbformat
from scipy.stats import binomtest
from scipy.stats.contingency import association
from docx import Document
from PIL import Image
import openpyxl
import pymupdf

WORK=Path(__file__).resolve().parents[1];BASE=WORK.parent
META=BASE/'data'/'metadata';TABLES=WORK/'tables'
config=json.loads((META/'dataset_config.json').read_text(encoding='utf-8'))
checks=[]
def check(name,condition):
    assert bool(condition),name
    checks.append(name)
def digest(path):
    h=hashlib.sha256()
    with path.open('rb') as f:
        for block in iter(lambda:f.read(1024*1024),b''):h.update(block)
    return h.hexdigest()
manifest=json.loads((META/'source_manifest.json').read_text(encoding='utf-8'))
for record in manifest:
    check('checksum '+record['file'],digest(BASE/record['file'])==record['sha256'])
check('local FRED equals official September vintage',manifest[-1]['matches_official_2026_09'])
loans=pd.read_parquet(BASE/'data/processed/loans_clean.parquet')
summary=json.loads((META/'eda_summary.json').read_text(encoding='utf-8'))
check('full sample count',len(loans)==summary['n_all'])
check('resolved label equivalence',loans.is_resolved.eq(loans.target.notna()).all())
check('target CHGOFF equals one',loans.loc[loans.LoanStatus.eq('CHGOFF'),'target'].eq(1).all())
check('target PIF equals zero',loans.loc[loans.LoanStatus.eq('PIF'),'target'].eq(0).all())
check('other statuses have no target',loans.loc[~loans.LoanStatus.isin(['PIF','CHGOFF']),'target'].isna().all())
check('snapshot exact',loans.AsOfDate.eq(pd.Timestamp(config['snapshot'])).all())
audit=json.loads((META/'pipeline_audit.json').read_text(encoding='utf-8'))
current_hashes={item['file']:digest(BASE/'data/raw'/item['file']) for item in config['sources']}
check('pipeline bound to active snapshot',audit.get('snapshot')==config['snapshot'])
check('pipeline bound to current source hashes',audit.get('source_hashes')==current_hashes)
for item in audit['files']:
    raw_counts=item['raw_status_counts']
    check('P I F normalization '+item['file'],int(loans.loc[loans.source_file.eq(item['file']),'LoanStatus'].eq('PIF').sum())==sum(v for k,v in raw_counts.items() if re.sub(r'\s+','',k.upper())=='PIF'))
    check('source dimensions '+item['file'],item['rows']==int(loans.source_file.eq(item['file']).sum()) and item['raw_columns']==len(item['raw_header']))
check('no invented Subprogram',not all('Subprogram' in item.get('unavailable_columns',[]) for item in audit['files']) or 'Subprogram' not in loans.columns)
check('FY range and date derivation',loans.fy.between(2000,2019).all() and loans.fy.eq(loans.ApprovalDate.dt.year+loans.ApprovalDate.dt.month.ge(10)).all())
deny={'borrname','borrstreet','borrcity','borrzip','locationid','bankfdicnumber','bankncuanumber'}
check('no borrower identifier fields in processed data',not deny.intersection(c.lower() for c in loans.columns))
for path in TABLES.glob('*.csv'):
    frame=pd.read_csv(path)
    check('no identifier columns '+path.name,not deny.intersection(c.lower() for c in frame.columns))
check('class counts',int(loans.target.eq(1).sum())==summary['n_chgoff'] and int(loans.target.eq(0).sum())==summary['n_pif'])
check('resolved rate',np.isclose(float(loans.target.mean()),summary['chgoff_rate']))
cohort=pd.read_csv(TABLES/'B8_cohort.csv')
check('cohort all reconciles',cohort.n_all.sum()==len(loans))
check('cohort resolved reconciles',cohort.n_resolved.sum()==loans.is_resolved.sum())
check('cohort CHGOFF reconciles',cohort.n_chgoff.sum()==loans.target.sum())

# Independent library reference checks: Wilson endpoints, including n=2,k=0.
states=pd.read_csv(TABLES/'B9_ty_le_BorrState.csv')
for row in states.loc[states.BorrState.isin(['AE','CA','FL','AK'])].itertuples():
    interval=binomtest(int(row.n_chgoff),int(row.n_resolved)).proportion_ci(method='wilson')
    check('Wilson '+row.BorrState,np.allclose([row.ci_low,row.ci_high],[interval.low,interval.high],rtol=1e-10,atol=1e-12))
ct=pd.crosstab(loans.BusinessType,loans.LoanStatus)
v=pd.read_csv(TABLES/'B16_cramers_v.csv').set_index('variable')
check('Cramers V agrees with SciPy',np.isclose(association(ct.to_numpy(),method='cramer',correction=False),v.loc['BusinessType','LoanStatus']))
join=json.loads((META/'macro_join_audit.json').read_text(encoding='utf-8'))
check('macro many-to-one preserves sample',join['n_before']==join['n_after']==len(loans))
check('macro months and missingness',join['matched_months']==240 and join['missing_lag_cells']==0)

doc=Document(WORK/'report/Chuong_2.docx')
text='\n'.join([p.text for p in doc.paragraphs]+[cell.text for t in doc.tables for row in t.rows for cell in row.cells])
md=(WORK/'report/Chuong_2.md').read_text(encoding='utf-8')
check('body font 13',doc.styles['Normal'].font.size.pt==13)
wb=openpyxl.load_workbook(BASE/'PTDLNC.xlsx',read_only=True,data_only=True)
assigned=[]
for row in wb.active.iter_rows(min_row=18,max_row=32,values_only=True):
    if row[0] is None:continue
    number=str(row[0]).strip().rstrip('.')
    title=str(row[1]).strip()
    if number=='2.1.1':title=title.replace('03/2026',config['snapshot_month_vi'])
    check('assigned heading '+number,number in text and title in text and title in md)
    assigned.append(number)
wb.close()

# Resolve report ledger back to original evidence, rather than trusting copied values.
ledger=pd.read_csv(TABLES/'Doi_chieu_so_lieu_bao_cao.csv')
for row in ledger.itertuples():
    if row.source.startswith('data/'):
        obj=json.loads((BASE/row.source).read_text(encoding='utf-8'))
        cursor=obj
        for key in re.findall(r'[^.\[\]]+',row.locator):
            cursor=cursor[int(key)] if isinstance(cursor,list) else cursor[key]
        expected=float(cursor)
    else:
        frame=pd.read_csv(WORK/row.source)
        key,column=row.locator.split('/')
        index='fy' if 'fy' in frame.columns and key.isdigit() else 'variable'
        expected=float(frame.loc[frame[index].astype(str).eq(key),column].iloc[0])
        if column=='missing_pct':expected/=100
    check('fact value '+row.fact,np.isclose(expected,row.value,rtol=1e-11,atol=1e-12))
    check('fact in Word '+row.fact,str(row.display) in text)
    check('fact in Markdown '+row.fact,str(row.display) in md)
check('at least ten independent report facts',len(ledger)>=10)

for path in (WORK/'notebooks').glob('*.ipynb'):
    nb=nbformat.read(path,as_version=4);nbformat.validate(nb)
    code=[c for c in nb.cells if c.cell_type=='code']
    check('all cells executed '+path.name,all(c.execution_count is not None for c in code))
    check('no cell errors '+path.name,not any(o.output_type=='error' for c in code for o in c.get('outputs',[])))
    # No borrower fields should occur as rendered table column headers.
    html='\n'.join(str(o.get('data',{}).get('text/html','')) for c in code for o in c.get('outputs',[]))
    check('no borrower headers in notebook outputs '+path.name,not any(re.search(r'<th[^>]*>\s*'+c+r'\s*</th>',html,re.I) for c in deny))

pngs=list((WORK/'figures').glob('*.png'))
for path in pngs:
    with Image.open(path) as img:
        check('figure DPI '+path.name,abs(img.info.get('dpi',(0,0))[0]-200)<1)
        check('figure size '+path.name,img.width>=1000 and img.height>=600)
image_paths=re.findall(r'!\[[^\]]*\]\(([^)]+)\)',md)
check('all Markdown image links resolve',all((WORK/'report'/link).exists() for link in image_paths))
check('eight report illustrations',len(image_paths)==8 and len(doc.inline_shapes)==8)
check('same number of tables',len(doc.tables)==md.count('**Bảng 2.'))
check('no old vintage in active report','260331' not in md and '31/03/2026' not in md)
check('active vintage in reference URLs',config['vintage'] in md)

pdf=WORK/'report/Chuong_2.pdf'
pagination={}
if pdf.exists():
    with pymupdf.open(pdf) as pages:
        reference_page=next(i+1 for i,p in enumerate(pages) if 'Tài liệu tham khảo' in p.get_text())
        pagination={'main_pages':reference_page-1,'total_pages':len(pages),'tool':'Microsoft Word export; inspected with PyMuPDF'}
        check('eight main pages',reference_page-1==8)
        check('no near-empty main pages',all(len(p.get_text())>350 for p in list(pages)[:reference_page-1]))
report={'result':'passed','checks_passed':len(checks),'checks':checks,'report_facts_checked':len(ledger),
        'snapshot':config['snapshot'],'raw_source_hashes':current_hashes,
        'assigned_sections':assigned,'pagination':pagination,'colab_execution':'not_verified_in_Colab',
        'local_notebooks':'passed','notes':'All requested statistical work is descriptive; no models trained.'}
(META/'deliverable_validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
build=json.loads((META/'report_build.json').read_text(encoding='utf-8'));build['actual_pagination']=pagination
(META/'report_build.json').write_text(json.dumps(build,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in report.items() if k!='checks'},ensure_ascii=False,indent=2))
